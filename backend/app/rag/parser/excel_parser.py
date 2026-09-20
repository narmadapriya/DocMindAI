"""XLSX parser: worksheets, tables, embedded images and chart metadata."""

from __future__ import annotations

from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Sequence

from app.rag.parser.exceptions import DocumentExtractionError
from app.rag.parser.models import (
    ChartBlock,
    DocumentMetadata,
    ImageBlock,
    ParsedDocument,
    TableBlock,
    TextBlock,
)


_MAX_RECORDS_PER_BLOCK = 12
_MAX_RECORD_BLOCK_CHARS = 2600
_MONTH_NAMES = {
    "jan": "January",
    "feb": "February",
    "mar": "March",
    "apr": "April",
    "may": "May",
    "jun": "June",
    "jul": "July",
    "aug": "August",
    "sep": "September",
    "sept": "September",
    "oct": "October",
    "nov": "November",
    "dec": "December",
}


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _retrieval_value(value: str, label: str = "") -> str:
    cleaned = str(value or "").strip()
    full_month = _MONTH_NAMES.get(cleaned.lower())
    if full_month and full_month.lower() != cleaned.lower():
        return f"{cleaned} ({full_month})"

    if "revenue" in str(label).lower():
        numeric = cleaned.replace("$", "").replace(",", "").strip()
        try:
            amount = Decimal(numeric)
        except InvalidOperation:
            amount = None

        if amount is not None and abs(amount) >= Decimal("1000000"):
            millions = (amount / Decimal("1000000")).normalize()
            million_text = format(millions, "f").rstrip("0").rstrip(".")
            return f"{cleaned} ({million_text} million USD)"

    return cleaned


def _normalize_headers(headers: Sequence[str], width: int) -> list[str]:
    result: list[str] = []
    for index in range(width):
        value = headers[index].strip() if index < len(headers) else ""
        result.append(value or f"Column {index + 1}")
    return result


def _non_empty_count(values: Sequence[str]) -> int:
    return sum(1 for value in values if str(value).strip())


def _is_header_row(values: Sequence[str], next_values: Sequence[str] | None) -> bool:
    non_empty = [str(value).strip() for value in values if str(value).strip()]
    if len(non_empty) < 2:
        return False
    if not all(not value.replace(".", "", 1).isdigit() for value in non_empty):
        return False
    if next_values is None:
        return False
    return _non_empty_count(next_values) >= 2




def _is_summary_row(values: Sequence[str]) -> bool:
    first = next((str(value).strip().lower() for value in values if str(value).strip()), "")
    return any(token in first for token in ("total", "subtotal", "average", "avg"))


def _serialize_record(headers: Sequence[str], values: Sequence[str], row_number: int) -> str:
    width = max(len(headers), len(values))
    labels = _normalize_headers(headers, width)
    row = list(values) + [""] * max(0, width - len(values))
    fields = [
        f"{labels[index]}: {_retrieval_value(row[index], labels[index])}"
        for index in range(width)
        if str(row[index]).strip()
    ]
    return f"Row {row_number}: " + "; ".join(fields)


def _worksheet_record_blocks(
    sheet_name: str,
    numbered_rows: list[tuple[int, list[str]]],
) -> list[TextBlock]:
    """
    Build bounded, retrieval-oriented worksheet row groups while preserving the
    original worksheet TableBlock.

    Very small three-row groups fragment ordinary spreadsheet questions across
    several evidence chunks.  This is especially expensive with a local Qwen
    verifier/retry loop.  Keep up to twelve logically adjacent records together
    when the serialized text remains bounded; large worksheets still split by
    section, blank-row gaps, summary rows, row count and character budget.
    """

    blocks: list[TextBlock] = []
    current_headers: list[str] | None = None
    current_context: list[str] = []
    batch: list[tuple[int, list[str]]] = []
    batch_chars = 0
    previous_row_number: int | None = None

    def flush_batch() -> None:
        nonlocal batch, batch_chars
        if not batch:
            return

        first_row = batch[0][0]
        last_row = batch[-1][0]
        row_range = f"{first_row}-{last_row}"
        width = max((len(values) for _, values in batch), default=0)
        headers = current_headers or [f"Column {index + 1}" for index in range(width)]

        lines = [f"Worksheet '{sheet_name}'."]
        if current_context:
            lines.append("Context: " + " | ".join(current_context[-2:]))
        lines.append("Columns: " + ", ".join(_normalize_headers(headers, max(len(headers), width))))
        lines.append(f"Rows {row_range}:")
        lines.extend(
            _serialize_record(headers, values, row_number)
            for row_number, values in batch
        )

        blocks.append(
            TextBlock(
                text="\n".join(lines),
                section=current_context[-1] if current_context else sheet_name,
                order=len(blocks),
                metadata={
                    "sheet_name": sheet_name,
                    "row_range": row_range,
                    "source": "xlsx_records",
                    "record_count": len(batch),
                },
            )
        )
        batch = []
        batch_chars = 0

    for index, (row_number, values) in enumerate(numbered_rows):
        next_values = numbered_rows[index + 1][1] if index + 1 < len(numbered_rows) else None

        # A blank-row gap normally separates logical worksheet sections/tables.
        if previous_row_number is not None and row_number - previous_row_number > 1:
            flush_batch()
            if current_headers is not None:
                current_headers = None
                current_context = []
        previous_row_number = row_number

        non_empty = [value for value in values if str(value).strip()]

        if len(non_empty) == 1:
            flush_batch()
            current_headers = None
            context_value = str(non_empty[0]).strip()
            current_context.append(context_value)
            current_context = current_context[-2:]
            blocks.append(
                TextBlock(
                    text=f"Worksheet '{sheet_name}'.\n{context_value}",
                    section=context_value,
                    order=len(blocks),
                    metadata={
                        "sheet_name": sheet_name,
                        "row_range": str(row_number),
                        "source": "xlsx_context",
                    },
                )
            )
            continue

        if _is_header_row(values, next_values):
            flush_batch()
            current_headers = list(values)
            continue

        if _is_summary_row(values):
            flush_batch()
            batch.append((row_number, values))
            batch_chars = len(_serialize_record(current_headers or [], values, row_number))
            flush_batch()
            continue

        headers_for_size = current_headers or []
        record_text = _serialize_record(headers_for_size, values, row_number)
        projected = batch_chars + len(record_text) + (1 if batch else 0)

        if batch and (
            len(batch) >= _MAX_RECORDS_PER_BLOCK
            or projected > _MAX_RECORD_BLOCK_CHARS
        ):
            flush_batch()

        batch.append((row_number, values))
        batch_chars += len(record_text) + (1 if batch_chars else 0)

    flush_batch()
    return blocks

def _resolved_row_values(value_sheet: Any, formula_sheet: Any) -> list[tuple[int, list[str]]]:
    rows: list[tuple[int, list[str]]] = []
    max_row = max(value_sheet.max_row, formula_sheet.max_row)
    max_column = max(value_sheet.max_column, formula_sheet.max_column)

    for row_number in range(1, max_row + 1):
        values: list[str] = []
        for column_number in range(1, max_column + 1):
            cached = value_sheet.cell(row=row_number, column=column_number).value
            formula = formula_sheet.cell(row=row_number, column=column_number).value
            resolved = cached if cached is not None else formula
            values.append(_cell_to_text(resolved))

        if any(values):
            rows.append((row_number, values))

    return rows


def parse_xlsx(file_path: str | Path, output_dir: str | Path | None = None) -> ParsedDocument:
    path = Path(file_path)
    if not path.exists():
        raise DocumentExtractionError(f"XLSX not found: {path}")

    try:
        import openpyxl
    except ImportError as exc:
        raise DocumentExtractionError("openpyxl is required for XLSX parsing.") from exc

    output = Path(output_dir) if output_dir else path.parent / "parsed_assets" / path.stem
    output.mkdir(parents=True, exist_ok=True)

    value_workbook = None
    formula_workbook = None

    try:
        # Use cached/evaluated formula values for retrieval.  Keep a second
        # formula workbook only as a fallback when a workbook has no cached value.
        value_workbook = openpyxl.load_workbook(path, data_only=True)
        formula_workbook = openpyxl.load_workbook(path, data_only=False)

        text_blocks: list[TextBlock] = []
        tables: list[TableBlock] = []
        images: list[ImageBlock] = []
        charts: list[ChartBlock] = []

        for value_sheet in value_workbook.worksheets:
            formula_sheet = formula_workbook[value_sheet.title]
            numbered_rows = _resolved_row_values(value_sheet, formula_sheet)
            rows = [values for _, values in numbered_rows]

            if rows:
                tables.append(
                    TableBlock(
                        rows=rows,
                        headers=rows[0] if len(rows) > 1 else None,
                        sheet_name=value_sheet.title,
                        table_id=f"xlsx-{value_sheet.title}-table1",
                        order=len(tables),
                        metadata={
                            "source": "xlsx_worksheet",
                            "max_row": value_sheet.max_row,
                            "max_column": value_sheet.max_column,
                            "formula_values": "cached_with_formula_fallback",
                        },
                    )
                )

                text_blocks.append(
                    TextBlock(
                        text=(
                            f"Worksheet '{value_sheet.title}' contains {len(rows)} non-empty rows "
                            f"and {max((len(row) for row in rows), default=0)} columns."
                        ),
                        section=value_sheet.title,
                        order=len(text_blocks),
                        metadata={
                            "sheet_name": value_sheet.title,
                            "source": "xlsx_worksheet_summary",
                        },
                    )
                )

                row_blocks = _worksheet_record_blocks(value_sheet.title, numbered_rows)
                for block in row_blocks:
                    block.order = len(text_blocks)
                    text_blocks.append(block)

            for image_index, image in enumerate(getattr(value_sheet, "_images", []), start=1):
                try:
                    image_bytes = image._data()
                    extension = getattr(image, "format", None) or "png"
                    image_path = output / f"{value_sheet.title}_image_{image_index}.{extension}"
                    image_path.write_bytes(image_bytes)
                    anchor_from = getattr(getattr(image, "anchor", None), "_from", None)
                    images.append(
                        ImageBlock(
                            path=str(image_path),
                            sheet_name=value_sheet.title,
                            image_id=f"xlsx-{value_sheet.title}-img{image_index}",
                            mime_type=f"image/{extension}",
                            width=getattr(image, "width", None),
                            height=getattr(image, "height", None),
                            order=len(images),
                            metadata={
                                "anchor": getattr(anchor_from, "col", None),
                                "needs_vision": True,
                                "source": "xlsx_embedded_image",
                            },
                        )
                    )
                except Exception as image_exc:
                    text_blocks.append(
                        TextBlock(
                            text=(
                                f"Unable to extract image {image_index} from worksheet "
                                f"'{value_sheet.title}': {image_exc}"
                            ),
                            section=value_sheet.title,
                            order=len(text_blocks),
                            metadata={
                                "sheet_name": value_sheet.title,
                                "source": "xlsx_image_error",
                            },
                        )
                    )

            for chart_index, chart in enumerate(getattr(value_sheet, "_charts", []), start=1):
                title = None
                try:
                    title = str(chart.title.tx.rich.p[0].r[0].t) if chart.title and chart.title.tx.rich else None
                except Exception:
                    title = None
                anchor_from = getattr(getattr(chart, "anchor", None), "_from", None)
                charts.append(
                    ChartBlock(
                        chart_id=f"xlsx-{value_sheet.title}-chart{chart_index}",
                        sheet_name=value_sheet.title,
                        chart_type=chart.__class__.__name__,
                        title=title,
                        order=len(charts),
                        metadata={
                            "anchor": getattr(anchor_from, "col", None),
                            "needs_visual_rendering": True,
                            "needs_vision": True,
                            "source": "xlsx_chart",
                        },
                    )
                )

        metadata = DocumentMetadata(
            filename=path.name,
            file_type=".xlsx",
            file_size=path.stat().st_size,
            sheet_names=value_workbook.sheetnames,
            extra={"sheet_count": len(value_workbook.sheetnames)},
        )

        return ParsedDocument(
            metadata=metadata,
            text_blocks=text_blocks,
            tables=tables,
            images=images,
            charts=charts,
        )
    except DocumentExtractionError:
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to parse XLSX '{path.name}': {exc}") from exc
    finally:
        if value_workbook is not None:
            value_workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
